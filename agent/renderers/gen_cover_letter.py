"""
Render a cover-letter dict (see agent/resume_tailor.py) to a .docx that shares
the resume's name banner and contact line.

Usage:
    python gen_cover_letter.py <input.json> <output.docx>

Input JSON shape:
{
  "name": "JORDAN RIVERA",
  "contact": "Denver, CO • jordan@example.com • linkedin.com/in/example",
  "date": "April 26, 2026",
  "recipient": "Hiring Team, Example Automation",
  "paragraphs": ["First paragraph...", "Second paragraph..."],
  "closing": "Best,",
  "signature": "Jordan Rivera"
}
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _add_run(p, text, *, bold=False, italic=False, size=None, font="Calibri"):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = font
    if size is not None:
        r.font.size = Pt(size)
    return r


def _set_space(p, before=0, after=0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def build(data):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # letterhead
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_space(p, before=0, after=2)
    _add_run(p, data["name"], bold=True, size=20)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_space(p, before=0, after=14)
    _add_run(p, data["contact"], size=10)

    # date
    p = doc.add_paragraph()
    _set_space(p, before=0, after=10)
    _add_run(p, data["date"], size=11)

    # recipient
    p = doc.add_paragraph()
    _set_space(p, before=0, after=10)
    _add_run(p, data["recipient"], size=11)

    # body paragraphs
    for para in data["paragraphs"]:
        p = doc.add_paragraph()
        _set_space(p, before=0, after=8)
        _add_run(p, para, size=11)

    # closing
    p = doc.add_paragraph()
    _set_space(p, before=8, after=2)
    _add_run(p, data.get("closing", "Best,"), size=11)

    p = doc.add_paragraph()
    _set_space(p, before=0, after=0)
    _add_run(p, data.get("signature", data["name"].title()), size=11)

    return doc


def main():
    if len(sys.argv) != 3:
        print("Usage: python gen_cover_letter.py <input.json> <output.docx>", file=sys.stderr)
        sys.exit(2)
    in_path, out_path = sys.argv[1], sys.argv[2]
    data = json.loads(Path(in_path).read_text(encoding="utf-8"))
    doc = build(data)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"wrote {out_path}")


if __name__ == "__main__":
    main()
