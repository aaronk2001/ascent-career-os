"""
Render a resume dict (see agent/resume_tailor.py) to a one-page .docx.

Usage:
    python gen_resume.py <input.json> <output.docx>

The input JSON has the shape:
{
  "name": "JORDAN RIVERA",
  "contact": "Denver, CO • jordan@example.com • linkedin.com/in/example",
  "tagline": "Controls Engineer | Automation • Computer Vision • Embedded Systems",
  "summary": "Controls engineer with ...",
  "experience": [
    {
      "company": "Example Robotics",
      "role": "Controls Engineer",
      "location": "Denver, CO",
      "dates": "Jan 2023 – Present",
      "bullets": [{"lead": "Cut changeover time 40%", "rest": " by ..."}]
    }
  ],
  "featured_project": {
      "header": "Capstone — Example University",
      "role": "Project Lead",
      "location": "Denver, CO",
      "dates": "Jan 2022 – May 2022",
      "bullets": [{"lead": "Built a 6-DOF arm ...", "rest": " ..."}]
  },
  "skills": [{"label": "Languages:", "content": "Python, C/C++, Ladder Logic"}],
  "education": "B.S., Mechanical Engineering — Example University",
  "education_date": "May 2022"
}
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ---------- low-level helpers ----------

def _set_paragraph_space(p, before=0, after=0, line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line


def _add_bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")        # 1pt
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_run(p, text, *, bold=False, italic=False, size=None, font="Calibri"):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = font
    if size is not None:
        r.font.size = Pt(size)
    return r


# ---------- section builders ----------

def add_header_block(doc, name, contact, tagline):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_space(p, before=0, after=2)
    _add_run(p, name, bold=True, size=20)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_space(p, before=0, after=2)
    _add_run(p, contact, size=10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_space(p, before=0, after=6)
    _add_run(p, tagline, italic=True, size=10)


def add_section_header(doc, label):
    p = doc.add_paragraph()
    _set_paragraph_space(p, before=8, after=2)
    _add_run(p, label.upper(), bold=True, size=11)
    _add_bottom_border(p)


def add_summary(doc, text):
    p = doc.add_paragraph()
    _set_paragraph_space(p, before=4, after=4)
    _add_run(p, text, size=10.5)


def add_role_heading(doc, *, header_left, role, location, dates):
    """Render a job heading like:  '<bold>Company  |  Role</bold>'  ............ '<italic>Location | Dates</italic>'"""
    p = doc.add_paragraph()
    _set_paragraph_space(p, before=4, after=2)
    # right-align tab at the right margin
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)

    _add_run(p, f"{header_left}  |  ", bold=True, size=10.5)
    _add_run(p, role, bold=True, size=10.5)
    _add_run(p, "\t", size=10.5)
    _add_run(p, f"{location}  |  {dates}", italic=True, size=10.5)


def add_company_heading(doc, *, company, location, note=None, roles=None):
    """v2 layout: '<bold>Company</bold>' .... '<bold>Location</bold>', an optional
    italic company description line, then one bold role/dates row per title."""
    p = doc.add_paragraph()
    _set_paragraph_space(p, before=4, after=0)
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
    _add_run(p, company, bold=True, size=10.5)
    _add_run(p, "\t", size=10.5)
    _add_run(p, location, bold=True, size=10.5)

    if note:
        p = doc.add_paragraph()
        _set_paragraph_space(p, before=0, after=0)
        _add_run(p, note, italic=True, size=9.5)

    for r in roles or []:
        p = doc.add_paragraph()
        _set_paragraph_space(p, before=0, after=0)
        pf = p.paragraph_format
        pf.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
        _add_run(p, r["title"], bold=True, size=10.5)
        _add_run(p, "\t", size=10.5)
        _add_run(p, r["dates"], bold=True, size=10.5)


def add_bullet(doc, lead, rest):
    p = doc.add_paragraph(style="List Bullet")
    _set_paragraph_space(p, before=0, after=2)
    pf = p.paragraph_format
    pf.left_indent = Inches(0.25)
    if lead:
        _add_run(p, lead, bold=True, size=10.5)
    if rest:
        _add_run(p, rest, size=10.5)


def add_skill_row(doc, label, content):
    p = doc.add_paragraph()
    _set_paragraph_space(p, before=2, after=2)
    _add_run(p, f"{label} ", bold=True, size=10.5)
    _add_run(p, content, size=10.5)


def add_education(doc, line, dates):
    p = doc.add_paragraph()
    _set_paragraph_space(p, before=4, after=2)
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Inches(7.0), WD_TAB_ALIGNMENT.RIGHT)
    # split first chunk (degree) bold, rest normal
    if " — " in line:
        degree, school = line.split(" — ", 1)
        _add_run(p, degree, bold=True, size=10.5)
        _add_run(p, " — " + school, size=10.5)
    else:
        _add_run(p, line, bold=True, size=10.5)
    _add_run(p, "\t", size=10.5)
    _add_run(p, dates, italic=True, size=10.5)


# ---------- main ----------

def build(data):
    doc = Document()

    # tight margins for one-page layout
    for section in doc.sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # default style
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    add_header_block(doc, data["name"], data["contact"], data["tagline"])

    add_section_header(doc, "Professional Summary")
    add_summary(doc, data["summary"])

    add_section_header(doc, "Professional Experience")
    for job in data["experience"]:
        if job.get("roles"):
            # v2 layout: company + note line + stacked role/date rows
            add_company_heading(
                doc,
                company=job["company"],
                location=job["location"],
                note=job.get("company_note"),
                roles=job["roles"],
            )
        else:
            add_role_heading(
                doc,
                header_left=job["company"],
                role=job["role"],
                location=job["location"],
                dates=job["dates"],
            )
        for b in job["bullets"]:
            add_bullet(doc, b.get("lead", ""), b.get("rest", ""))

    if data.get("featured_project"):
        fp = data["featured_project"]
        add_section_header(doc, "Featured Project")
        add_role_heading(
            doc,
            header_left=fp["header"],
            role=fp["role"],
            location=fp["location"],
            dates=fp["dates"],
        )
        for b in fp["bullets"]:
            add_bullet(doc, b.get("lead", ""), b.get("rest", ""))

    add_section_header(doc, "Technical Skills")
    for s in data["skills"]:
        add_skill_row(doc, s["label"], s["content"])

    add_section_header(doc, "Education")
    add_education(doc, data["education"], data.get("education_date", ""))

    return doc


def main():
    if len(sys.argv) != 3:
        print("Usage: python gen_resume.py <input.json> <output.docx>", file=sys.stderr)
        sys.exit(2)
    in_path, out_path = sys.argv[1], sys.argv[2]
    data = json.loads(Path(in_path).read_text(encoding="utf-8"))
    doc = build(data)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"wrote {out_path}")


if __name__ == "__main__":
    main()
